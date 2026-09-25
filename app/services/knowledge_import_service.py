import logging
import re
import unicodedata
import zipfile
from dataclasses import dataclass
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.knowledge_entry import INTENT_LABELS, KNOWLEDGE_INTENTS, KnowledgeEntry, SourceType
from app.services.errors import BusinessError
from app.services.text_service import URL_PATTERN, normalize_major, normalize_text

logger = logging.getLogger(__name__)

TEMPLATE_FILENAME = "Mau_Import_TriThuc.xlsx"
GUIDE_SHEET = "Hướng dẫn"
EXAMPLE_MARK = "VÍ DỤ"
XLSX_MAGIC = b"PK\x03\x04"
NOT_XLSX = "Chỉ nhận tệp Excel .xlsx. Hãy tải file mẫu, điền dữ liệu rồi lưu lại dưới dạng .xlsx."
MEGABYTE = 1024 * 1024
MAJOR_MAX_LENGTH = 150
TITLE_MAX_LENGTH = 255
CONTENT_MAX_LENGTH = 60000
URL_MAX_LENGTH = 1000
YEAR_RANGE = range(2000, 2101)
HEADER_NOISE = re.compile(r"\(.*?\)|\*")


@dataclass(frozen=True)
class Column:
    key: str
    label: str
    required: bool
    width: int
    guide: str

    @property
    def header(self) -> str:
        return f"{self.label} *" if self.required else self.label


IMPORT_COLUMNS = (
    Column("intent", "Chủ đề (intent)", True, 26, "Mã chủ đề ở bảng bên dưới, ví dụ hoi_hoc_phi."),
    Column("major", "Ngành học", False, 28, "Tên ngành. Để trống nếu áp dụng cho mọi ngành."),
    Column("year", "Năm", False, 8, "Năm áp dụng, 4 chữ số, ví dụ 2025. Để trống nếu không theo năm."),
    Column("title", "Tiêu đề", True, 42, "Tối đa 255 ký tự."),
    Column("content", "Nội dung", True, 70, "Câu trả lời. Ghi đúng số liệu, điều khoản trường công bố."),
    Column("source_url", "Nguồn (URL)", False, 40, "Đường dẫn http/https tới trang gốc (không bắt buộc)."),
)
COLUMN_LABELS = {column.key: column.label for column in IMPORT_COLUMNS}


@dataclass(frozen=True)
class ImportRow:
    row: int
    intent: str
    major: str | None
    year: int | None
    title: str
    content: str
    source_url: str | None


@dataclass(frozen=True)
class RowError:
    row: int
    column: str
    message: str


@dataclass(frozen=True)
class ParsedImport:
    rows: list[ImportRow]
    errors: list[RowError]
    total_rows: int


@dataclass(frozen=True)
class ImportResult:
    total_rows: int
    errors: list[RowError]
    message: str


def _header_key(value) -> str:
    return normalize_text(HEADER_NOISE.sub("", str(value or ""))).lower()


HEADER_KEYS = {_header_key(column.label): column.key for column in IMPORT_COLUMNS}


def build_template() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Tri thức"
    sheet.append([column.header for column in IMPORT_COLUMNS])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0E6E78")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for index, column in enumerate(IMPORT_COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = column.width
    sheet.append(["hoi_diem_chuan", "Hệ thống thông tin", 2025, f"{EXAMPLE_MARK}: Điểm chuẩn ngành Hệ thống thông tin 2025",
                  "Ghi điểm chuẩn đúng theo công bố của trường cho từng phương thức xét tuyển.", None])
    sheet.append(["hoi_quy_che_hoc_vu", None, None, f"{EXAMPLE_MARK}: Điều kiện xét tốt nghiệp",
                  "Chép nguyên văn điều khoản trong quy chế học vụ.", None])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    guide = workbook.create_sheet(GUIDE_SHEET)
    guide.append(["Cột", "Bắt buộc", "Cách điền"])
    for column in IMPORT_COLUMNS:
        guide.append([column.label, "Có" if column.required else "Không", column.guide])
    guide.append([])
    guide.append(["Mã chủ đề (intent)", "Ý nghĩa"])
    for intent in KNOWLEDGE_INTENTS:
        guide.append([intent.value, INTENT_LABELS[intent]])
    guide.append([])
    guide.append([f"Dòng có chữ “{EXAMPLE_MARK}” được bỏ qua khi nhập. Nhập lại cùng chủ đề, ngành, năm, tiêu đề sẽ cập nhật nội dung, không tạo bản trùng."])
    for cell in (*guide[1], *guide[len(IMPORT_COLUMNS) + 3]):
        cell.font = Font(bold=True)
    guide.column_dimensions["A"].width = 28
    guide.column_dimensions["B"].width = 34
    guide.column_dimensions["C"].width = 70

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _text(value) -> str:
    return normalize_text(str(value)) if value is not None else ""


def _year(value) -> int | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool) and float(value).is_integer():
        return int(value)
    if isinstance(value, str) and value.strip().isdigit():
        return int(value.strip())
    raise ValueError


def _parse_row(excel_row: int, cells: dict) -> tuple[ImportRow | None, list[RowError]]:
    errors: list[RowError] = []

    def fail(key: str, message: str) -> None:
        errors.append(RowError(excel_row, COLUMN_LABELS[key], message))

    raw_intent = _text(cells.get("intent"))
    intent = raw_intent.lower()
    if intent not in KNOWLEDGE_INTENTS:
        fail("intent", f"Chủ đề “{raw_intent}” không hợp lệ. Dùng một mã trong sheet Hướng dẫn." if raw_intent
             else "Thiếu chủ đề. Dùng một mã trong sheet Hướng dẫn.")

    major = _text(cells.get("major")) or None
    if major and len(major) > MAJOR_MAX_LENGTH:
        fail("major", f"Tên ngành dài quá {MAJOR_MAX_LENGTH} ký tự.")

    year = None
    try:
        year = _year(cells.get("year"))
    except ValueError:
        fail("year", "Năm phải là số nguyên 4 chữ số, ví dụ 2025.")
    if year is not None and year not in YEAR_RANGE:
        fail("year", f"Năm {year} nằm ngoài khoảng {YEAR_RANGE.start}–{YEAR_RANGE.stop - 1}.")

    title = _text(cells.get("title"))
    if not title:
        fail("title", "Thiếu tiêu đề.")
    elif len(title) > TITLE_MAX_LENGTH:
        fail("title", f"Tiêu đề dài quá {TITLE_MAX_LENGTH} ký tự.")

    raw_content = cells.get("content")
    content = unicodedata.normalize("NFC", str(raw_content)).strip() if raw_content is not None else ""
    if not content:
        fail("content", "Thiếu nội dung.")
    elif len(content) > CONTENT_MAX_LENGTH:
        fail("content", f"Nội dung dài quá {CONTENT_MAX_LENGTH} ký tự. Hãy tách thành nhiều dòng.")

    source_url = _text(cells.get("source_url")) or None
    if source_url and (not URL_PATTERN.match(source_url) or len(source_url) > URL_MAX_LENGTH):
        fail("source_url", "Nguồn phải là đường dẫn bắt đầu bằng http:// hoặc https://.")

    if errors:
        return None, errors
    return ImportRow(excel_row, intent, major, year, title, content, source_url), []


def parse_workbook(data: bytes) -> ParsedImport:
    if not data.startswith(XLSX_MAGIC):
        raise BusinessError(NOT_XLSX)
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError) as error:
        raise BusinessError(NOT_XLSX) from error

    try:
        values = workbook.worksheets[0].iter_rows(values_only=True)
        positions: dict[str, int] = {}
        for index, header in enumerate(next(values, None) or ()):
            key = HEADER_KEYS.get(_header_key(header))
            if key and key not in positions:
                positions[key] = index
        missing = [column.label for column in IMPORT_COLUMNS if column.required and column.key not in positions]
        if missing:
            raise BusinessError(f"Tệp thiếu cột bắt buộc: {', '.join(missing)}. Hãy dùng đúng file mẫu.")

        rows, errors, total = [], [], 0
        for excel_row, row_values in enumerate(values, start=2):
            cells = {key: row_values[index] if index < len(row_values) else None for key, index in positions.items()}
            texts = [_text(value) for value in cells.values()]
            if not any(texts) or any(EXAMPLE_MARK in text.upper() for text in texts):
                continue
            total += 1
            parsed, row_errors = _parse_row(excel_row, cells)
            errors.extend(row_errors)
            if parsed:
                rows.append(parsed)
    finally:
        workbook.close()
    return ParsedImport(rows=rows, errors=errors, total_rows=total)


def _find_existing(db: Session, row: ImportRow, major_norm: str | None) -> KnowledgeEntry | None:
    return db.scalar(select(KnowledgeEntry).where(
        KnowledgeEntry.source_type == SourceType.EXCEL,
        KnowledgeEntry.intent == row.intent,
        KnowledgeEntry.title == row.title,
        KnowledgeEntry.major_norm.is_not_distinct_from(major_norm),
        KnowledgeEntry.year.is_not_distinct_from(row.year),
    ))


def import_file(db: Session, data: bytes, user_id: int, max_mb: int) -> ImportResult:
    if len(data) > max_mb * MEGABYTE:
        raise BusinessError(f"Tệp vượt quá {max_mb} MB. Hãy tách thành nhiều tệp nhỏ hơn.")
    parsed = parse_workbook(data)
    if parsed.total_rows == 0:
        raise BusinessError(f"Tệp không có dòng dữ liệu nào (các dòng “{EXAMPLE_MARK}” được bỏ qua).")
    if parsed.errors:
        bad_rows = len({error.row for error in parsed.errors})
        message = f"Chưa nhập dòng nào: {bad_rows}/{parsed.total_rows} dòng có lỗi. Sửa các lỗi bên dưới rồi tải lại tệp."
        return ImportResult(parsed.total_rows, parsed.errors, message)

    inserted = updated = 0
    for row in parsed.rows:
        major_norm = normalize_major(row.major)
        existing = _find_existing(db, row, major_norm)
        if existing:
            existing.content, existing.major, existing.source_url = row.content, row.major, row.source_url
            updated += 1
            continue
        db.add(KnowledgeEntry(
            intent=row.intent, major=row.major, major_norm=major_norm, year=row.year, title=row.title,
            content=row.content, source_type=SourceType.EXCEL, source_url=row.source_url, created_by=user_id,
        ))
        inserted += 1
    db.commit()
    logger.info("Import tri thức Excel: %s dòng (thêm %s, cập nhật %s) bởi user %s", parsed.total_rows, inserted, updated, user_id)
    message = f"Đã nhập {parsed.total_rows} dòng tri thức: thêm mới {inserted}, cập nhật {updated}."
    return ImportResult(parsed.total_rows, [], message)
